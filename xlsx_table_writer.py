"""Write a DataFrame out as a real Excel Table (ListObject), not a plain range.

The original workbooks all used Excel Tables (e.g. "Odoo_2026_05"), and the
embedded Power Query logic filters specifically on Kind = "Table" when
scanning the folder - a plain pandas.to_excel() range would be invisible to
that logic and just looks like an unstructured sheet to anyone opening it.
"""

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def write_as_table(df, path, table_name, sheet_name="Odoo Report"):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        ws = writer.sheets[sheet_name]

        n_rows, n_cols = df.shape
        last_col = get_column_letter(n_cols)
        ref = f"A1:{last_col}{n_rows + 1}"

        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight9", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)
